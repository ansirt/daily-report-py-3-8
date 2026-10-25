from __future__ import annotations

import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Border, Font, Side
from openpyxl.utils import get_column_letter

_THIN = Side(style="thin", color="000000")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_OK = (str, int, float, bool, Decimal, datetime, date, time, timedelta)


def clean(v):
    if v is None:
        return None
    if isinstance(v, (bytes, bytearray, memoryview)):
        return bytes(v).hex()
    if not isinstance(v, _OK):
        v = str(v)
    if isinstance(v, datetime) and v.tzinfo:
        v = v.replace(tzinfo=None)
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def write_table(ws, cols, rows, top=1, left=1):
    for j, name in enumerate(cols):
        c = ws.cell(top, left + j, name)
        c.font = Font(bold=True)
        c.border = BORDER
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            c = ws.cell(top + i, left + j, clean(v))
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"  # data, bukan formula
            c.border = BORDER


def autofit(ws, skip_rows=(), sample=500):
    widths = {}
    for row in ws.iter_rows(min_row=1, max_row=sample + 1):
        for c in row:
            if c.value is None or c.row in skip_rows:
                continue
            n = len(str(c.value))
            widths[c.column] = max(widths.get(c.column, 0), n)
    for col, n in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(60, max(8, n + 2))


def _sheet_title(name, used):
    t = re.sub(r"[\[\]:*?/\\]", "_", name or "Sheet")[:31] or "Sheet"
    base, i = t, 2
    while t.lower() in used:
        t = f"{base[:28]}_{i}"
        i += 1
    used.add(t.lower())
    return t


def build_standard(path, sheets):
    """sheets: list of (sheet_name, columns, rows)"""
    wb = Workbook(write_only=True)
    wb.remove(wb.active)
    used = set()
    for name, cols, rows in sheets or [("Sheet1", [], [])]:
        ws = wb.create_sheet(_sheet_title(name, used))
        write_table(ws, cols, rows)
        autofit(ws)
    wb.save(path)


# ---- layout summary_order (port dari summaryOrder.cs) ----
_DRY = [
    ("Order Gms", '=SUMIFS(Sheet2!$B:$B,Sheet2!$C:$C,"GMS")'),
    ("Order Gro", '=SUMIFS(Sheet2!$B:$B,Sheet2!$C:$C,"GRO")'),
    ("Order Shoppingbag", '=SUMIFS(Sheet2!$B:$B,Sheet2!$C:$C,"SB")'),
    ("Obcarton Packed", '=SUMIFS(Sheet2!$B:$B,Sheet2!$A:$A,"Packed")'),
    ("Obcarton Printed", '=SUMIFS(Sheet2!$B:$B,Sheet2!$A:$A,"Printed")'),
    ("Total collie plan picking", '=SUMIFS(Sheet3!$D:$D,Sheet3!$B:$B,"D")'),
    ("Total Truck rencana muat", '=COUNTIFS(Sheet3!$B:$B,"D")-COUNTIFS(Sheet3!$B:$B,"D",Sheet3!$C:$C,"CONTAINER*")'),
    ("Total Container rencana muat", '=COUNTIFS(Sheet3!$B:$B,"D",Sheet3!$C:$C,"CONTAINER*")'),
    ("Total booking Reguler", '=SUMIFS(Sheet4!$E:$E,Sheet4!$F:$F,"REG",Sheet4!$G:$G,"DRY")'),
    ("Total booking Xdk", '=SUMIFS(Sheet4!$E:$E,Sheet4!$F:$F,"XDK",Sheet4!$G:$G,"DRY")'),
    ("Total Po Booking", '=COUNTIF(Sheet4!$G:$G,"DRY")'),
    ("Order yang masuk hari H & H-1 - Release", '=SUMIFS(Sheet2!$B:$B,Sheet2!$C:$C,"All Order H and H-1",Sheet2!$A:$A,"Released")'),
    ("Order yang masuk hari H & H-1 NEW AUTOREP - Release", '=SUMIFS(Sheet2!$B:$B,Sheet2!$C:$C,"All Order H and H-1 NEW AUTOREP :autorep",Sheet2!$A:$A,"Released")'),
    ("Total order yang masuk hari H & H-1", '=SUMIFS(Sheet2!$B:$B,Sheet2!$C:$C,"All Order H and H-1")'),
    ("Total order yang masuk hari H & H-1 NEW AUTOREP", '=SUMIFS(Sheet2!$B:$B,Sheet2!$C:$C,"All Order H and H-1 NEW AUTOREP :autorep")'),
]
_FRESH = [
    ("Order Frozen", '=SUMIFS(Sheet2!$F:$F,Sheet2!$G:$G,"FRZ")'),
    ("Order Chiller", '=SUMIFS(Sheet2!$F:$F,Sheet2!$G:$G,"CHI")'),
    ("Obcarton Packed", '=SUMIFS(Sheet2!$F:$F,Sheet2!$E:$E,"Packed")'),
    ("Obcarton Printed", '=SUMIFS(Sheet2!$F:$F,Sheet2!$E:$E,"Printed")'),
    ("Total collie plan picking", '=SUMIFS(Sheet3!$D:$D,Sheet3!$B:$B,"F")'),
    ("Total Truck rencana muat", '=COUNTIFS(Sheet3!$B:$B,"F")-COUNTIFS(Sheet3!$B:$B,"F",Sheet3!$C:$C,"CONTAINER*")'),
    ("Total Container rencana muat", '=COUNTIFS(Sheet3!$B:$B,"F",Sheet3!$C:$C,"CONTAINER*")'),
    ("Total booking Reguler", '=SUMIFS(Sheet4!$E:$E,Sheet4!$F:$F,"REG",Sheet4!$G:$G,"<>DRY")'),
    ("Total booking Xdk", '=SUMIFS(Sheet4!$E:$E,Sheet4!$F:$F,"XDK",Sheet4!$G:$G,"<>DRY")'),
    ("Total Po Booking", '=COUNTIF(Sheet4!$G:$G,"FRS")'),
    ("Order yang masuk hari H & H-1 - Release", '=SUMIFS(Sheet2!$F:$F,Sheet2!$G:$G,"All Order H and H-1",Sheet2!$E:$E,"Released")'),
    ("Order yang masuk hari H & H-1 NEW AUTOREP - Release", '=SUMIFS(Sheet2!$F:$F,Sheet2!$G:$G,"All Order H and H-1 NEW AUTOREP :autorep",Sheet2!$E:$E,"Released")'),
    ("Total order yang masuk hari H & H-1", '=SUMIFS(Sheet2!$F:$F,Sheet2!$G:$G,"All Order H and H-1")'),
    ("Order yang masuk hari H & H-1 NEW AUTOREP", '=SUMIFS(Sheet2!$F:$F,Sheet2!$G:$G,"All Order H and H-1 NEW AUTOREP :autorep")'),
]


def build_summary_order(path, res, title_date):
    """res: {posisi_query: (cols, rows)} -> 0=dry, 1=fresh, 2=plan delivery, 3=po booking."""
    empty = ([], [])
    wb = Workbook(write_only=True)
    s1 = wb.active
    s1.title = "Summary"
    s2, s3, s4 = wb.create_sheet("Sheet2"), wb.create_sheet("Sheet3"), wb.create_sheet("Sheet4")
    s2.cell(1, 1, "Dry")
    write_table(s2, *res.get(0, empty), top=2, left=1)
    s2.cell(1, 5, "Fresh")
    write_table(s2, *res.get(1, empty), top=2, left=5)
    write_table(s3, *res.get(2, empty))
    write_table(s4, *res.get(3, empty))
    for ws in (s2, s3, s4):
        autofit(ws, skip_rows=(1,) if ws is s2 else ())

    s1.cell(1, 1, f"Berikut Data Summary Order DC Porong,  {title_date}")
    s1.cell(3, 1, "Dry"); s1.cell(3, 2, "Value")
    s1.cell(3, 4, "Fresh"); s1.cell(3, 5, "Value")
    for i, (label, f) in enumerate(_DRY):
        s1.cell(4 + i, 1, label); s1.cell(4 + i, 2, f)
    for i, (label, f) in enumerate(_FRESH):
        s1.cell(4 + i, 4, label); s1.cell(4 + i, 5, f)
    last = s1.max_row
    for r in range(3, last + 1):
        for col in (1, 2, 4, 5):
            s1.cell(r, col).border = BORDER
    autofit(s1, skip_rows=(1,))
    wb.active = 0
    wb.save(path)
